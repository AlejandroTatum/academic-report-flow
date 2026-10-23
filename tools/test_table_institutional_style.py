"""Tests for the institutional UNL table style TAB-IN-08.

TAB-IN-08 is an explicit-override-only style: it has a UNL-blue header
(#1F5F8B, white bold text), soft alternating rows and left alignment, and
it must never be picked by automatic contextual selection.
"""
from __future__ import annotations

import itertools
import sys
from pathlib import Path

import pytest
import yaml

TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS))

from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from table_docx_tokens import render_styled_table_docx  # noqa: E402
from table_html_tokens import render_styled_table_html  # noqa: E402
from table_latex_tokens import render_styled_table_latex  # noqa: E402
from table_model import TableRequest, resolve_table_style  # noqa: E402
from table_styles import (  # noqa: E402
    APPROVED_STYLE_IDS,
    HEADER_FILL_HEX,
    CatalogError,
    TableContext,
    eligible_candidates,
    load_catalog,
    select_style,
)
from test_table_styles import COMPARISON, DENSE, LONG, MULTIPAGE, SHORT, STATUS  # noqa: E402

ROOT = TOOLS.parent
STYLE_ID = "TAB-IN-08"
CATALOG = load_catalog()
STYLE = CATALOG.styles[STYLE_ID]

HEADER = ["Item", "Valor"]
ROWS = [["A", "1"], ["B", "2"], ["C", "3"]]


def test_tab_in_08_is_in_the_approved_catalog_with_institutional_tokens() -> None:
    assert STYLE_ID in APPROVED_STYLE_IDS
    assert STYLE.label == "Institucional UNL"
    assert STYLE.deprecated is False
    assert STYLE.tokens == {
        "borders": "horizontal_only",
        "header": "institutional_shaded",
        "alignment": "left",
        "padding": "standard",
        "density": "medium",
        "row_rhythm": "alternating",
        "palette": "monochrome",
        "indicators": "none",
        "caption": "above",
        "notes": "below",
    }
    assert HEADER_FILL_HEX["institutional_shaded"] == "1F5F8B"


@pytest.mark.parametrize("context", [SHORT, LONG, COMPARISON, STATUS, DENSE, MULTIPAGE])
def test_tab_in_08_override_resolves_for_any_context(context: TableContext) -> None:
    request = TableRequest(table_key="tabla-1", context=context, teacher_override=STYLE_ID)
    receipt = resolve_table_style(request, CATALOG)
    assert receipt.style_id == STYLE_ID
    assert receipt.precedence_source == "teacher"


def test_tab_in_08_institution_override_resolves() -> None:
    request = TableRequest(table_key="tabla-1", context=SHORT, institution_override=STYLE_ID)
    assert resolve_table_style(request, CATALOG).style_id == STYLE_ID


def test_automatic_selection_results_are_unchanged_and_never_tab_in_08() -> None:
    expected = {
        "short": ("TAB-CL-01", SHORT),
        "long": ("TAB-ZB-04", LONG),
        "comparison": ("TAB-CE-05", COMPARISON),
        "status": ("TAB-ES-06", STATUS),
        "dense": ("TAB-CC-07", DENSE),
        "multipage": ("TAB-ZB-04", MULTIPAGE),
    }
    for name, (expected_id, context) in expected.items():
        assert select_style(context, CATALOG).style_id == expected_id, name
        assert STYLE_ID not in eligible_candidates(context, CATALOG), name


def test_automatic_selection_never_returns_tab_in_08_over_the_whole_context_space() -> None:
    """Exhaustive sweep: no context makes TAB-IN-08 a candidate or a winner,
    and removing it from the catalog changes no outcome (the 'existing
    reports render byte-identically' guarantee, at the selector level)."""
    without = type(CATALOG)(
        schema_version=CATALOG.schema_version,
        catalog_version=CATALOG.catalog_version,
        styles={k: v for k, v in CATALOG.styles.items() if k != STYLE_ID},
        status_indicators=CATALOG.status_indicators,
        source_path=CATALOG.source_path,
    )
    axes = itertools.product(
        ("reference", "comparison", "status", "evidence", "dense"),
        ("short", "long"),
        ("low", "medium", "high"),
        (1, 3, 4, 8),
        ("none", "column"),
        ("none", "comparison", "status"),
        ("color", "grayscale"),
    )
    for purpose, length, density, columns, emphasis, meaning, color_policy in axes:
        context = TableContext(
            purpose=purpose, length=length, density=density, columns=columns, rows=10,
            emphasis=emphasis, meaning=meaning, color_policy=color_policy,
        )
        assert STYLE_ID not in eligible_candidates(context, CATALOG)
        assert eligible_candidates(context, CATALOG) == eligible_candidates(context, without)


def test_automatic_key_must_be_a_bool(tmp_path: Path) -> None:
    raw = yaml.safe_load((ROOT / "templates" / "table_styles.yml").read_text(encoding="utf-8"))
    raw["styles"][STYLE_ID]["automatic"] = "no"
    path = tmp_path / "table_styles.yml"
    path.write_text(yaml.safe_dump(raw), encoding="utf-8")
    with pytest.raises(CatalogError, match="'automatic' must be a bool"):
        load_catalog(path)


def test_tab_in_08_latex_renders_blue_header_white_text_left_aligned() -> None:
    text = "\n".join(render_styled_table_latex(
        header=HEADER, rows=ROWS, style=STYLE, status_indicators=CATALOG.status_indicators,
        convert_inline=lambda value: value,
    ))
    assert f"% table-style: {STYLE_ID}" in text
    assert r"\rowcolor[HTML]{1F5F8B}" in text
    assert r"{\color{white} \textbf{Item}}" in text
    assert r"\raggedright" in text
    assert r"\rowcolors" in text


def test_tab_in_08_html_renders_blue_header_white_text_left_aligned() -> None:
    html = render_styled_table_html(
        header=HEADER, rows=ROWS, style=STYLE, status_indicators=CATALOG.status_indicators,
        escape_inline=lambda value: value,
    )
    assert f'data-table-style="{STYLE_ID}"' in html
    assert "background:#1f5f8b;color:#ffffff;" in html
    assert "text-align:left" in html
    assert "#f2f2f2" in html


def test_tab_in_08_docx_renders_blue_header_white_text_left_aligned() -> None:
    def fill_cell(cell, text: str, *, bold: bool, size_pt: float) -> None:
        run = cell.paragraphs[0].add_run(text)
        run.bold = bold

    table = render_styled_table_docx(
        document=Document(), header=HEADER, rows=ROWS, style=STYLE,
        status_indicators=CATALOG.status_indicators, fill_cell=fill_cell,
    )
    for cell in table.rows[0].cells:
        assert cell._tc.get_or_add_tcPr().find(qn("w:shd")).get(qn("w:fill")) == "1F5F8B"
        runs = cell.paragraphs[0].runs
        assert runs and all(str(run.font.color.rgb) == "FFFFFF" and run.bold for run in runs)
    assert table.rows[1].cells[0].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert table.rows[2].cells[0]._tc.get_or_add_tcPr().find(qn("w:shd")).get(qn("w:fill")) == "F2F2F2"
